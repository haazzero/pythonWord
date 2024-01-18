# Word 라이브러리
import docx
from docx import Document
# Excel 라이브러리
import openpyxl
from openpyxl import Workbook
# 파일 로컬에서 열기
import win32com.client

# excel 어플리케이션 열기
excel = win32com.client.gencache.EnsureDispatch("Excel.application")
# 자동으로 실행되는 것을 확인하고 싶을 때 코드 추가
excel.Visible = True

# 워드 파일 경로
word_file_path = 'C:\\Users\\hy\\OneDrive\\eng\\Short Essay_hy.docx'

# 워드 파일 열기
doc = Document(word_file_path)

# 형광펜 단어를 담을 공간
highlighted_text = []

# 단락 읽기와 font 객체
for para in doc.paragraphs:
    for run in para.runs :
        if run.font.highlight_color:
            highlighted_text.append(run.text)

print(f'모르는 단어 수 : {len(highlighted_text)}')
print(highlighted_text)

#####################

# Workbook 객체 생성
wb = Workbook()
# 활성화된 워크시트 선택 후 ws 변수에 할당
ws = wb.active
# 시트 제목을 Highlighted Word로 설정
ws.title = "Highlighted Words" 

# 첫 행에 1열과 2열에 단어, 뜻을 입력
first = ['단어','뜻']
ws.append(first)

# highlighted_text에 담긴 단어를 for문으로 'A'열의 각 셀에 순서대로 접근하여 입력
for i, value in enumerate(highlighted_text):
    ws.cell(row=i+2, column=1, value=value)

# 엑셀 파일 경로 및 제목 설정
excel_file_path = word_file_path.replace('.docx', '_voca.xlsx')
# 워크북 엑셀 파일 저장 (제목)
wb.save(excel_file_path)

# 저장된 엑셀 파일 열기
excel.Workbooks.Open(excel_file_path)
print(f'저장 된 엑셀 파일 경로 : {excel_file_path}')

# 워크북 닫기
wb.close()