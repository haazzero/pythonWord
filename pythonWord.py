import docx
from docx import Document
import openpyxl
from openpyxl import Workbook
import win32com.client
import requests
from bs4 import BeautifulSoup

# Excel 어플리케이션 열기
excel = win32com.client.gencache.EnsureDispatch("Excel.application")
excel.Visible = True

# 워드 파일 경로 입력받기
word_file_path = input('워드 파일 경로를 입력해주세요.').strip('"')

# 워드 파일 열기
doc = Document(word_file_path)

# 형광펜 단어를 담을 공간
highlighted_text = []
# 형광펜 단어를 포함한 문장을 담을 공간
highlighted_text_sentence = []
# 형광펜 단어의 한국어 뜻을 담을 공간
highlighted_text_korean = []

# 단락 읽기와 font 객체
# 문장을 '.'을 기준으로 나누고 각 문장에 대해 형광펜 단어 추출
for para in doc.paragraphs:
    sentences = para.text.split('. ')
    for sentence in sentences:
        for run in para.runs:
            if run.font.highlight_color and run.text in sentence:
                highlighted_text.append(run.text)
                highlighted_text_sentence.append(sentence)

# 크롤링하여 한글 뜻 가져오기
for word in highlighted_text:
    url = f"https://www.wordreference.com/enko/{word}"
    response = requests.get(url)
    if response.status_code == 200:
        html = response.text
        soup = BeautifulSoup(html, 'html.parser')
        result = soup.select_one('tr.even > td.ToWrd')
        korean = result.get_text()
        highlighted_text_korean.append(korean)

# 결과 출력
print('*************** 형광펜 단어에 대한 정보를 모았어요! ***************')
print(f'*************** 모르는 단어 수: {len(highlighted_text)}')
for i in range(len(highlighted_text)):
    print(f"단어: {highlighted_text[i]}")
    print(f"문장: {highlighted_text_sentence[i]}")
    print(f"뜻: {highlighted_text_korean[i]}")

# Excel 파일 생성 및 데이터 입력
wb = Workbook()
ws = wb.active
ws.title = "Highlighted Words" 

header = ['번호', '단어', '뜻', '문장']
ws.append(header)

for i in range(len(highlighted_text)):
    word = highlighted_text[i]
    sentence = highlighted_text_sentence[i]
    mean = highlighted_text_korean[i]
    ws.append([i+1, word, mean, sentence])

excel_file_path = word_file_path.replace('.docx', '_voca.xlsx')
wb.save(excel_file_path)

# 저장된 엑셀 파일 열기
excel.Workbooks.Open(excel_file_path)
print(f'저장된 엑셀 파일 경로: {excel_file_path}')

# 워크북 닫기
wb.close()
